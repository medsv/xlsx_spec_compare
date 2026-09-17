# app.py
# Python >= 3.14
#
# Сравнение двух спецификаций Excel.
# Результат — XLSX с листами:
#   1) Итог
#   2) Изменения
#   3) Добавленные
#   4) Удалённые
#   5) Спецификация (разметка) — копия листа из r1 с сохранением оформления,
#      добавленными удалёнными строками из r0 и цветовой разметкой.

import io
import re
import unicodedata
from collections import defaultdict
from copy import copy
from dataclasses import dataclass, field
from datetime import datetime
from typing import Any

import pandas as pd
import streamlit as st
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


HEADER_SEARCH_LIMIT = 30

PLACEHOLDER_RE = re.compile(r"^\d+[.)]?$")

HOMOGLYPH_MAP = {
    "А": "A", "а": "a",
    "В": "B", "в": "b",
    "Е": "E", "е": "e",
    "К": "K", "к": "k",
    "М": "M", "м": "m",
    "Н": "H", "н": "h",
    "О": "O", "о": "o",
    "Р": "P", "р": "p",
    "С": "C", "с": "c",
    "Т": "T", "т": "t",
    "Х": "X", "х": "x",
    "У": "Y", "у": "y",
}

HOMOGLYPH_TRANSLATE = {ord(k): v for k, v in HOMOGLYPH_MAP.items()}


@dataclass(slots=True)
class FieldChange:
    field: str
    old: str
    new: str


@dataclass(slots=True)
class RowRecord:
    excel_row: int
    section: str
    is_section: bool
    key: str
    values: dict[str, str]


@dataclass(slots=True)
class ParsedSpec:
    file_name: str
    title: str
    header_row: int
    columns: list[str]
    records: list[RowRecord]


@dataclass(slots=True)
class RowChange:
    key: str
    section: str
    old_row_num: int | None = None
    new_row_num: int | None = None
    old_row: dict[str, str] | None = None
    new_row: dict[str, str] | None = None
    changed_fields: list[FieldChange] = field(default_factory=list)


@dataclass(slots=True)
class ComparisonResult:
    meta: dict[str, Any]
    summary: dict[str, int]
    columns: list[str]
    field_map: dict[str, str | None]
    added: list[RowChange]
    deleted: list[RowChange]
    modified: list[RowChange]


def normalize_text(value: Any) -> str:
    if value is None:
        return ""

    try:
        if pd.isna(value):
            return ""
    except TypeError:
        pass

    if isinstance(value, bool):
        return str(value)

    if isinstance(value, (int, float)):
        try:
            f = float(value)
            if f.is_integer():
                return str(int(f))
        except Exception:
            pass
        return str(value)

    s = str(value)
    s = unicodedata.normalize("NFKC", s)
    s = s.replace("\xa0", " ")
    s = re.sub(r"\s+", " ", s).strip()

    if s in {"-", "—", "–", "--", "---"}:
        return ""

    return s


def normalize_for_key(value: Any) -> str:
    s = normalize_text(value)
    if not s:
        return ""

    s = s.translate(HOMOGLYPH_TRANSLATE)

    s = (
        s.replace('"', "")
        .replace("«", "")
        .replace("»", "")
        .replace("'", "")
        .replace("’", "")
    )

    s = re.sub(r"\s+", " ", s).strip().lower()
    return s


def clean_header(value: Any, idx: int) -> str:
    result = normalize_text(value)
    return result if result else f"Столбец {idx + 1}"


def make_unique_columns(columns: list[str]) -> list[str]:
    seen: dict[str, int] = {}
    result: list[str] = []

    for col in columns:
        if col not in seen:
            seen[col] = 1
            result.append(col)
        else:
            seen[col] += 1
            result.append(f"{col} ({seen[col]})")

    return result


def find_column(columns: list[str], keywords: tuple[str, ...]) -> str | None:
    for keyword in keywords:
        for col in columns:
            if keyword in col.lower():
                return col
    return None


def build_field_map(columns: list[str]) -> dict[str, str | None]:
    return {
        "poz": find_column(columns, ("поз. по спецификации", "поз")),
        "name": find_column(columns, ("наименование",)),
        "type": find_column(columns, ("тип,", "тип")),
        "code": find_column(columns, ("код оборудования", "код")),
        "factory": find_column(columns, ("завод",)),
        "unit": find_column(columns, ("ед. изм", "ед")),
        "qty": find_column(columns, ("кол-во", "количество")),
        "mass": find_column(columns, ("масса",)),
        "note": find_column(columns, ("примечание",)),
    }


def read_title(file_bytes: bytes) -> str:
    wb = load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=True)
    ws = wb.active
    title = ws.cell(row=1, column=1).value if ws is not None else ""
    wb.close()
    return normalize_text(title)


def find_header_index(df_raw: pd.DataFrame) -> int:
    limit = min(len(df_raw), HEADER_SEARCH_LIMIT)

    for i in range(limit):
        row = df_raw.iloc[i]
        texts = [normalize_text(v).lower() for v in row]
        joined = " | ".join(texts)

        if "поз" in joined and "наименование" in joined:
            return i

    for i in range(limit):
        row = df_raw.iloc[i]
        non_empty = sum(1 for v in row if normalize_text(v))
        if non_empty >= 3:
            return i

    return 1 if len(df_raw) > 1 else 0


def parse_spec(file_bytes: bytes, file_name: str) -> ParsedSpec:
    title = read_title(file_bytes)

    df_raw = pd.read_excel(
        io.BytesIO(file_bytes),
        header=None,
        dtype=object,
        engine="openpyxl",
        sheet_name=0,
    )

    if df_raw.empty:
        return ParsedSpec(
            file_name=file_name,
            title=title,
            header_row=1,
            columns=[],
            records=[],
        )

    header_idx = find_header_index(df_raw)
    header_excel_row = header_idx + 1

    header_values = df_raw.iloc[header_idx]

    if len(header_values) == 0:
        return ParsedSpec(
            file_name=file_name,
            title=title,
            header_row=header_excel_row,
            columns=[],
            records=[],
        )

    last_header_col = -1
    for i, v in enumerate(header_values):
        if normalize_text(v):
            last_header_col = i

    if last_header_col < 0:
        last_header_col = len(header_values) - 1

    raw_headers = header_values.iloc[: last_header_col + 1].tolist()
    columns = make_unique_columns(
        [clean_header(v, i) for i, v in enumerate(raw_headers)]
    )

    field_map = build_field_map(columns)

    poz_col = field_map["poz"] or (columns[0] if columns else None)
    name_col = field_map["name"] or (columns[1] if len(columns) > 1 else None)
    type_col = field_map["type"]
    code_col = field_map["code"]
    factory_col = field_map["factory"]
    unit_col = field_map["unit"]
    qty_col = field_map["qty"]
    mass_col = field_map["mass"]
    note_col = field_map["note"]

    volatile_cols = {c for c in (qty_col, mass_col, note_col) if c}

    data = df_raw.iloc[header_idx + 1 :, : last_header_col + 1].reset_index(drop=True)

    records: list[RowRecord] = []

    current_section = ""
    current_section_key = ""

    section_counts: dict[str, int] = {}
    key_counts: dict[str, int] = {}

    for offset, row in data.iterrows():
        excel_row = header_excel_row + offset + 1

        values: dict[str, str] = {}

        for i, col in enumerate(columns):
            values[col] = normalize_text(row[i]) if i < len(row) else ""

        non_empty = [col for col, val in values.items() if val]

        if not non_empty:
            continue

        if len(non_empty) == 1:
            only_value = values[non_empty[0]]
            if PLACEHOLDER_RE.fullmatch(only_value):
                continue

        is_section = bool(name_col) and len(non_empty) == 1 and non_empty[0] == name_col

        if is_section:
            section_name = values[name_col]
            section_key = normalize_for_key(section_name)

            if section_key == "госты":
                break

            count = section_counts.get(section_key, 0) + 1
            section_counts[section_key] = count

            current_section = (
                section_name if count == 1 else f"{section_name} ({count})"
            )
            current_section_key = (
                section_key if count == 1 else f"{section_key}#{count}"
            )

            key = f"section::{current_section_key}"

        else:
            poz = values.get(poz_col, "") if poz_col else ""

            name_val = values.get(name_col, "") if name_col else ""
            unit_val = values.get(unit_col, "") if unit_col else ""
            qty_val = values.get(qty_col, "") if qty_col else ""

            is_item = bool(name_val and (unit_val or qty_val))

            if is_item:
                item_parts: list[str] = []

                for col in (name_col, type_col, code_col, factory_col, unit_col):
                    if col:
                        item_parts.append(normalize_for_key(values.get(col, "")))

                base_key = (
                    f"item::{current_section_key or 'no_section'}::"
                    + "|".join(item_parts)
                )

            elif poz:
                base_key = f"poz::{normalize_for_key(poz)}"

            else:
                fallback_parts: list[str] = []

                for col, val in values.items():
                    if col in volatile_cols:
                        continue

                    if val:
                        fallback_parts.append(
                            f"{normalize_for_key(col)}={normalize_for_key(val)}"
                        )

                if fallback_parts:
                    base_key = "row::" + "|".join(fallback_parts)
                else:
                    base_key = f"empty::{offset}"

            count = key_counts.get(base_key, 0) + 1
            key_counts[base_key] = count

            key = base_key if count == 1 else f"{base_key}#{count}"

        records.append(
            RowRecord(
                excel_row=excel_row,
                section=current_section,
                is_section=is_section,
                key=key,
                values=values,
            )
        )

    return ParsedSpec(
        file_name=file_name,
        title=title,
        header_row=header_excel_row,
        columns=columns,
        records=records,
    )


def compare_specs(old: ParsedSpec, new: ParsedSpec) -> ComparisonResult:
    columns = list(dict.fromkeys(old.columns + new.columns))
    field_map = build_field_map(columns)

    old_map = {record.key: record for record in old.records}
    new_map = {record.key: record for record in new.records}

    old_keys = [record.key for record in old.records]
    new_keys = [record.key for record in new.records]

    added_keys = [key for key in new_keys if key not in old_map]
    deleted_keys = [key for key in old_keys if key not in new_map]
    common_keys = [key for key in new_keys if key in old_map]

    added: list[RowChange] = []
    deleted: list[RowChange] = []
    modified: list[RowChange] = []
    unchanged = 0

    for key in added_keys:
        rec = new_map[key]
        added.append(
            RowChange(
                key=key,
                section=rec.section,
                new_row_num=rec.excel_row,
                new_row=rec.values,
            )
        )

    for key in deleted_keys:
        rec = old_map[key]
        deleted.append(
            RowChange(
                key=key,
                section=rec.section,
                old_row_num=rec.excel_row,
                old_row=rec.values,
            )
        )

    for key in common_keys:
        old_rec = old_map[key]
        new_rec = new_map[key]

        changed_fields: list[FieldChange] = []

        for col in columns:
            old_val = old_rec.values.get(col, "")
            new_val = new_rec.values.get(col, "")

            if old_val != new_val:
                changed_fields.append(
                    FieldChange(
                        field=col,
                        old=old_val,
                        new=new_val,
                    )
                )

        if changed_fields:
            modified.append(
                RowChange(
                    key=key,
                    section=new_rec.section,
                    old_row_num=old_rec.excel_row,
                    new_row_num=new_rec.excel_row,
                    old_row=old_rec.values,
                    new_row=new_rec.values,
                    changed_fields=changed_fields,
                )
            )
        else:
            unchanged += 1

    summary = {
        "added": len(added),
        "deleted": len(deleted),
        "modified": len(modified),
        "unchanged": unchanged,
        "total_old": len(old.records),
        "total_new": len(new.records),
    }

    meta = {
        "generated_at": datetime.now().isoformat(),
        "old_file": old.file_name,
        "new_file": new.file_name,
        "old_title": old.title,
        "new_title": new.title,
        "headers_equal": old.columns == new.columns,
    }

    return ComparisonResult(
        meta=meta,
        summary=summary,
        columns=columns,
        field_map=field_map,
        added=added,
        deleted=deleted,
        modified=modified,
    )


def changed_fields_to_df(result: ComparisonResult) -> pd.DataFrame:
    headers = [
        "Статус",
        "Раздел",
        "Позиция",
        "Наименование",
        "Изменённое поле",
        "Значение в r0",
        "Значение в r1",
        "Строка r0",
        "Строка r1",
    ]

    rows: list[dict[str, Any]] = []

    poz_col = result.field_map.get("poz")
    name_col = result.field_map.get("name")

    for item in result.modified:
        source = item.new_row or item.old_row or {}

        for fc in item.changed_fields:
            rows.append(
                {
                    "Статус": "Изменено",
                    "Раздел": item.section,
                    "Позиция": source.get(poz_col, "") if poz_col else "",
                    "Наименование": source.get(name_col, "") if name_col else "",
                    "Изменённое поле": fc.field,
                    "Значение в r0": fc.old,
                    "Значение в r1": fc.new,
                    "Строка r0": item.old_row_num,
                    "Строка r1": item.new_row_num,
                }
            )

    return pd.DataFrame(rows, columns=headers)


def items_to_df(
    items: list[RowChange],
    columns: list[str],
    mode: str,
) -> pd.DataFrame:
    headers = ["Раздел", "Excel row", *columns]
    rows: list[dict[str, Any]] = []

    for item in items:
        values = item.new_row if mode == "added" else item.old_row
        row_num = item.new_row_num if mode == "added" else item.old_row_num

        row: dict[str, Any] = {
            "Раздел": item.section,
            "Excel row": row_num,
        }

        values = values or {}

        for col in columns:
            row[col] = values.get(col, "")

        rows.append(row)

    return pd.DataFrame(rows, columns=headers)


def to_cell(value: Any) -> Any:
    if value is None:
        return ""

    try:
        if pd.isna(value):
            return ""
    except TypeError:
        pass

    return value


def df_to_rows(df: pd.DataFrame) -> list[list[Any]]:
    if df.empty:
        return []

    return [[to_cell(v) for v in row] for row in df.astype(object).values.tolist()]


def copy_row_style(ws, template_row: int, target_row: int, max_col: int) -> None:
    """
    Копирует стиль строки-образца в целевую строку.
    """
    if template_row < 1 or template_row > ws.max_row:
        return

    for col in range(1, max_col + 1):
        src = ws.cell(row=template_row, column=col)
        dst = ws.cell(row=target_row, column=col)

        dst.font = copy(src.font)
        dst.border = copy(src.border)
        dst.fill = copy(src.fill)
        dst.alignment = copy(src.alignment)
        dst.number_format = src.number_format
        dst.protection = copy(src.protection)

    src_dim = ws.row_dimensions.get(template_row)
    if src_dim is not None and src_dim.height:
        ws.row_dimensions[target_row].height = src_dim.height


def export_xlsx(
    result: ComparisonResult,
    old: ParsedSpec,
    new: ParsedSpec,
    new_file_bytes: bytes,
) -> bytes:
    """
    Создаёт итоговый файл:
    - листы с результатами сравнения;
    - лист-копию спецификации r1 с сохранением исходного оформления;
    - на этом листе добавляет удалённые строки и раскрашивает изменения.
    """
    # Загружаем исходный файл r1, чтобы сохранить оформление первого листа.
    wb_out = load_workbook(io.BytesIO(new_file_bytes))

    if not wb_out.worksheets:
        wb_out = Workbook()
        ws_marked = wb_out.active
        ws_marked.title = "Спецификация (разметка)"
    else:
        ws_marked = wb_out.worksheets[0]
        ws_marked.title = "Спецификация (разметка)"

        # Оставляем только лист спецификации, остальные листы r1 удаляем.
        for ws in list(wb_out.worksheets):
            if ws is not ws_marked:
                wb_out.remove(ws)

    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)

    added_fill = PatternFill("solid", fgColor="C6EFCE")
    deleted_fill = PatternFill("solid", fgColor="FFC7CE")
    modified_fill = PatternFill("solid", fgColor="FFEB9C")

    # Создаём служебные листы перед листом разметки.
    ws_summary = wb_out.create_sheet("Итог", 0)
    ws_changes = wb_out.create_sheet("Изменения", 1)
    ws_added = wb_out.create_sheet("Добавленные", 2)
    ws_deleted = wb_out.create_sheet("Удалённые", 3)

    def write_table(
        ws,
        headers: list[str],
        rows: list[list[Any]],
        fill: PatternFill | None = None,
    ) -> None:
        ws.append(headers)

        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(vertical="top", wrap_text=True)

        if rows:
            for row in rows:
                ws.append(row)

                if fill:
                    for cell in ws[ws.max_row]:
                        cell.fill = fill
        else:
            ws.append(["Нет данных"])

        for col_idx, header in enumerate(headers, start=1):
            max_len = len(str(header))
            max_row = min(ws.max_row, 100)

            for row in ws.iter_rows(
                min_row=2,
                max_row=max_row,
                min_col=col_idx,
                max_col=col_idx,
            ):
                for cell in row:
                    if cell.value is not None:
                        max_len = max(
                            max_len,
                            min(len(str(cell.value)), 80),
                        )

            ws.column_dimensions[get_column_letter(col_idx)].width = min(
                max_len + 2,
                80,
            )

        ws.freeze_panes = "A2"

    # Лист «Итог»
    summary_rows = [
        ["Файл до корректировки", result.meta.get("old_file", "")],
        ["Файл после корректировки", result.meta.get("new_file", "")],
        ["Название в r0", result.meta.get("old_title", "")],
        ["Название в r1", result.meta.get("new_title", "")],
        ["Дата сравнения", result.meta.get("generated_at", "")],
        ["Состав столбцов одинаковый", result.meta.get("headers_equal", "")],
        ["Добавлено строк", result.summary.get("added", 0)],
        ["Удалено строк", result.summary.get("deleted", 0)],
        ["Изменено строк", result.summary.get("modified", 0)],
        ["Без изменений", result.summary.get("unchanged", 0)],
        ["Всего строк в r0", result.summary.get("total_old", 0)],
        ["Всего строк в r1", result.summary.get("total_new", 0)],
    ]

    write_table(ws_summary, ["Параметр", "Значение"], summary_rows)

    # Лист «Изменения»
    changes_df = changed_fields_to_df(result)
    write_table(
        ws_changes,
        list(changes_df.columns),
        df_to_rows(changes_df),
        modified_fill,
    )

    # Лист «Добавленные»
    added_df = items_to_df(result.added, result.columns, "added")
    write_table(
        ws_added,
        list(added_df.columns),
        df_to_rows(added_df),
        added_fill,
    )

    # Лист «Удалённые»
    deleted_df = items_to_df(result.deleted, result.columns, "deleted")
    write_table(
        ws_deleted,
        list(deleted_df.columns),
        df_to_rows(deleted_df),
        deleted_fill,
    )

    # ==========================================================
    # Разметка листа-копии r1
    # ==========================================================

    marked_columns = new.columns if new.columns else result.columns
    marked_col_index = {col: i + 1 for i, col in enumerate(marked_columns)}

    max_col = max(ws_marked.max_column, len(marked_columns))

    deleted_keys = {item.key for item in result.deleted}
    added_keys = {item.key for item in result.added}
    modified_map = {item.key: item for item in result.modified}

    new_map = {rec.key: rec for rec in new.records}
    new_row_map = {rec.excel_row: rec for rec in new.records}

    old_keys_set = {rec.key for rec in old.records}
    new_keys_set = {rec.key for rec in new.records}
    common_keys_set = old_keys_set & new_keys_set

    # Для каждой удалённой строки ищем в r0 ближайшую общую строку,
    # которая шла после неё. В r1 удалённая строка будет вставлена
    # перед этой общей строкой.
    deleted_by_anchor: dict[int, list[RowRecord]] = defaultdict(list)
    pending_deleted: list[RowRecord] = []

    for rec in old.records:
        if rec.key in deleted_keys:
            pending_deleted.append(rec)
        elif rec.key in common_keys_set:
            if pending_deleted:
                anchor_rec = new_map.get(rec.key)
                if anchor_rec:
                    deleted_by_anchor[anchor_rec.excel_row].extend(pending_deleted)
                pending_deleted = []

    end_deleted = pending_deleted

    # Вставляем удалённые строки снизу вверх, чтобы не сбивать номера строк.
    for anchor_row in sorted(deleted_by_anchor.keys(), reverse=True):
        recs = deleted_by_anchor[anchor_row]
        if not recs:
            continue

        amount = len(recs)

        anchor_rec = new_row_map.get(anchor_row)

        # Если строка-якорь является разделом, обычно лучше брать стиль
        # из строки выше, а не из самого раздела.
        use_above_template = bool(
            anchor_rec
            and anchor_rec.is_section
            and (anchor_row - 1) > new.header_row
        )

        if anchor_row <= ws_marked.max_row:
            ws_marked.insert_rows(anchor_row, amount)

            if use_above_template:
                template_row = anchor_row - 1
            else:
                template_row = anchor_row + amount

                if template_row > ws_marked.max_row:
                    template_row = max(1, anchor_row - 1)
        else:
            template_row = max(1, anchor_row - 1)

        for i, rec in enumerate(recs):
            ins_row = anchor_row + i

            copy_row_style(ws_marked, template_row, ins_row, max_col)

            for col_idx, col in enumerate(marked_columns, start=1):
                ws_marked.cell(
                    row=ins_row,
                    column=col_idx,
                    value=to_cell(rec.values.get(col, "")),
                )

            for col_idx in range(1, max_col + 1):
                ws_marked.cell(row=ins_row, column=col_idx).fill = deleted_fill

    # Вставляем удалённые строки, которые в r0 шли после последней общей строки.
    if end_deleted:
        if new.records:
            last_parsed_row = new.records[-1].excel_row
        else:
            last_parsed_row = new.header_row

        total_shift = sum(
            len(recs)
            for anchor, recs in deleted_by_anchor.items()
            if anchor <= last_parsed_row
        )

        end_start = last_parsed_row + total_shift + 1
        amount = len(end_deleted)

        if end_start <= ws_marked.max_row:
            ws_marked.insert_rows(end_start, amount)
            template_row = max(1, end_start - 1)
        else:
            template_row = max(1, end_start - 1)

        for i, rec in enumerate(end_deleted):
            ins_row = end_start + i

            copy_row_style(ws_marked, template_row, ins_row, max_col)

            for col_idx, col in enumerate(marked_columns, start=1):
                ws_marked.cell(
                    row=ins_row,
                    column=col_idx,
                    value=to_cell(rec.values.get(col, "")),
                )

            for col_idx in range(1, max_col + 1):
                ws_marked.cell(row=ins_row, column=col_idx).fill = deleted_fill

    # После вставки удалённых строк вычисляем итоговые номера строк r1.
    final_row_by_key: dict[str, int] = {}

    for rec in new.records:
        shift = sum(
            len(recs)
            for anchor, recs in deleted_by_anchor.items()
            if anchor <= rec.excel_row
        )
        final_row_by_key[rec.key] = rec.excel_row + shift

    # Раскрашиваем добавленные строки.
    for key in added_keys:
        rec = new_map.get(key)
        if not rec:
            continue

        final_row = final_row_by_key.get(rec.key)
        if not final_row:
            continue

        for col_idx in range(1, max_col + 1):
            ws_marked.cell(row=final_row, column=col_idx).fill = added_fill

    # Раскрашиваем только изменённые ячейки.
    for item in result.modified:
        rec = new_map.get(item.key)
        if not rec:
            continue

        final_row = final_row_by_key.get(rec.key)
        if not final_row:
            continue

        for fc in item.changed_fields:
            col_idx = marked_col_index.get(fc.field)
            if col_idx:
                ws_marked.cell(row=final_row, column=col_idx).fill = modified_fill

    wb_out.active = 0

    bio = io.BytesIO()
    wb_out.save(bio)
    return bio.getvalue()


def main() -> None:
    st.set_page_config(
        page_title="Сравнение спецификаций",
        layout="wide",
    )

    st.title("Сравнение спецификаций")

    file_r0 = st.file_uploader("Файл r0", type=["xlsx"])
    file_r1 = st.file_uploader("Файл r1", type=["xlsx"])

    if st.button("Сравнить", type="primary"):
        if not file_r0 or not file_r1:
            st.warning("Нужно загрузить оба файла: r0 и r1.")
        else:
            try:
                old_bytes = file_r0.getvalue()
                new_bytes = file_r1.getvalue()

                old_spec = parse_spec(old_bytes, file_r0.name)
                new_spec = parse_spec(new_bytes, file_r1.name)

                result = compare_specs(old_spec, new_spec)

                st.session_state.result = result
                st.session_state.old_spec = old_spec
                st.session_state.new_spec = new_spec
                st.session_state.summary = result.summary
                st.session_state.xlsx_bytes = export_xlsx(
                    result,
                    old_spec,
                    new_spec,
                    new_bytes,
                )

            except Exception as exc:
                st.error(f"Ошибка при сравнении: {exc}")
                st.exception(exc)

    summary = st.session_state.get(
        "summary",
        {
            "added": 0,
            "deleted": 0,
            "modified": 0,
            "unchanged": 0,
        },
    )

    c1, c2, c3, c4 = st.columns(4)

    c1.metric("Добавлено", summary.get("added", 0))
    c2.metric("Удалено", summary.get("deleted", 0))
    c3.metric("Изменено", summary.get("modified", 0))
    c4.metric("Без изменений", summary.get("unchanged", 0))

    xlsx_bytes = st.session_state.get("xlsx_bytes")

    if xlsx_bytes:
        timestamp = datetime.now().strftime("%Y-%m-%d_%H-%M-%S")

        st.download_button(
            "Скачать результат сравнения",
            data=xlsx_bytes,
            file_name=f"spec_comparison_{timestamp}.xlsx",
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )


if __name__ == "__main__":
    main()
